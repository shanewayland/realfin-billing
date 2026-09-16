"""Monthly statement email: python testMonthly.py (no network, no real email)."""
import io, os, smtplib
from unittest import mock
from openpyxl import load_workbook
os.environ.update(CRON_KEY='k', BUBBLE_API_TOKEN='t', MAIL_FROM='from@example.com',
                  MAIL_TO='to@example.com', SMTP_PASSWORD='p')
import monthly
from service import app

LOANS = [
  {"_id": "L2", "loan_number": "002", "borrower_name": "CND-Everlight, LLC", "loan_amount": 11072238,
   "property_address": "Everlight S/D", "accrual start date": "2026-06-09T05:00:00.000Z",
   "advanced_at_closing": 158199, "loan_spread_rate": 0.035, "floor_rate": 0, "initial_prime_rate": 0.0675,
   "Status": "Closed"},
  {"_id": "L9", "loan_number": "009", "borrower_name": "No Start Date LLC", "advanced_at_closing": 0},
]
ACTS = [
  {"loan": "L2", "activity_date": "2026-07-01T05:00:00.000Z", "interest_payment": 990.94},
  {"loan": "L2", "activity_date": "2026-07-28T05:00:00.000Z", "disbursement_to_borrower": 758798.42},
  {"loan": "L2", "activity_date": "2026-08-01T05:00:00.000Z", "interest_payment": 2269.26},
  {"loan": "L2", "activity_date": "2026-09-01T05:00:00.000Z", "interest_payment": 8122.55},
  {"loan": "L2", "activity_date": "2026-09-03T05:00:00.000Z", "disbursement_to_borrower": 1223412.99},
]
P = F = 0
def check(label, ok, got=''):
    global P, F
    print(f"  {'PASS' if ok else 'FAIL'}  {label} {got}")
    P, F = (P + 1, F) if ok else (P, F + 1)

c = app.test_client()
H = {'X-Cron-Key': 'k'}
def fake_call(version, payload):
    if not payload:
        return {'loans': LOANS, 'activities': [], 'activity_count': 0}
    mine = [a for a in ACTS if a['loan'] == payload['loan']]
    off = payload['offset']
    return {'activities': mine[off:off + monthly.PAGE], 'activity_count': len(mine)}

with mock.patch.object(monthly, '_call_bubble', side_effect=fake_call):
    check('rejects missing key', c.post('/monthly-statements').status_code == 401)
    check('rejects wrong key', c.post('/monthly-statements', headers={'X-Cron-Key': 'x'}).status_code == 401)

    r = c.post('/monthly-statements?dry_run=1&period=2026-09-01', headers=H).json
    check('dry run builds 002', r['statements'] == ['002_2026-09_Billing_Statement.xlsx'], r['statements'])
    check('dry run lists skipped loan', len(r['skipped']) == 1 and '009' in r['skipped'][0], r['skipped'])
    check('dry run sends nothing', r['sent'] is False)

    sent = {}
    class FakeSMTP:
        def __init__(self, host, port, timeout=None): sent['host'] = (host, port)
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def starttls(self): sent['tls'] = True
        def login(self, u, p): sent['login'] = u
        def send_message(self, m): sent['msg'] = m
    with mock.patch.object(smtplib, 'SMTP', FakeSMTP):
        r = c.post('/monthly-statements?period=2026-09-01&to=test@example.com', headers=H).json
    m = sent['msg']
    check('sent', r['sent'] is True)
    check('to override', m['To'] == 'test@example.com', m['To'])
    check('from / login', m['From'] == 'from@example.com' and sent['login'] == 'from@example.com')
    check('tls on 587', sent['tls'] and sent['host'] == ('smtp.gmail.com', 587))
    check('subject', m['Subject'] == 'RealFin billing statements — September 2026', m['Subject'])
    atts = list(m.iter_attachments())
    check('one attachment', len(atts) == 1)
    ws = load_workbook(io.BytesIO(atts[0].get_payload(decode=True))).active
    total = [ws.cell(r, 8).value for r in range(1, 40) if ws.cell(r, 7).value == 'PLEASE PAY THIS AMOUNT:']
    check('Sep 2026 amount matches her statement ($17,683.23)', total == [17683.23], total)
    check('landscape layout kept', ws.page_setup.orientation == 'landscape')

    at = lambda *a: monthly.datetime(*a, tzinfo=monthly.CENTRAL)
    with mock.patch.object(monthly, 'now_central', return_value=at(2026, 10, 1, 9, 5)):
        r = c.post('/monthly-statements?scheduled=1', headers=H).json
        check('scheduled run skips at 9am (the second DST trigger)', 'skipped' in r, r)
    with mock.patch.object(monthly, 'now_central', return_value=at(2026, 10, 2, 8, 5)):
        r = c.post('/monthly-statements?scheduled=1', headers=H).json
        check('scheduled run skips on the 2nd', 'skipped' in r)
    with mock.patch.object(monthly, 'now_central', return_value=at(2026, 10, 1, 8, 2)):
        r = c.post('/monthly-statements?scheduled=1&dry_run=1', headers=H).json
        check('scheduled run at 8am uses last month', r.get('period') == 'September 2026', r.get('period'))

check('UTC midnight-CT date converts to Houston date', monthly.central_date('2026-09-01T05:00:00.000Z') == '2026-09-01')
check('prior month across year', monthly.prior_month(monthly.date(2027, 1, 1)) == monthly.date(2026, 12, 1))

# paging: 250 entries for one loan must all arrive (3 pages)
many = [{"loan": "L2", "activity_date": "2026-07-01T05:00:00.000Z", "interest_payment": 1}] * 250
calls = []
def paged(version, payload):
    calls.append(payload)
    off = payload['offset']
    return {'activities': many[off:off + 100], 'activity_count': 250}
with mock.patch.object(monthly, '_call_bubble', side_effect=paged):
    items, expected = monthly.fetch_activity('live', 'L2')
check('paging gets all 250 entries', len(items) == 250 and expected == 250, len(items))
check('paging used offsets 0/100/200', [c['offset'] for c in calls] == [0, 100, 200], [c['offset'] for c in calls])

# Bubble returned fewer entries than it counted: no statement for that loan
def short(version, payload):
    if not payload:
        return {'loans': LOANS[:1]}
    return {'activities': ACTS[:1], 'activity_count': 5}
with mock.patch.object(monthly, '_call_bubble', side_effect=short):
    r = c.post('/monthly-statements?dry_run=1&period=2026-09-01', headers=H).json
check('incomplete activity is skipped, not billed', r['statements'] == [] and 'only 1 of 5' in r['skipped'][0], r['skipped'])

print(f"\n  {P} passed, {F} failed")
